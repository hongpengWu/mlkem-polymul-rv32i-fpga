"""Rebuild the competition report with the bundled Python runtime.

Only reads project evidence. Does not invoke FPGA tools or board operations.
Usage: python docs/competition/build_report.py
       python docs/competition/build_report.py --template <other-template.docx>
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import os
import shutil
import zipfile
from functools import lru_cache

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ASSETS = HERE / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
BLUE = '#245477'
TEAL = '#087F8C'
INK = '#172B3A'
GRAY = '#566778'
PALE = '#EDF4F8'
GREEN = '#E8F4F2'
GOLD = '#AD7626'
SCALE = 2.4


@lru_cache(maxsize=2)
def report_font_path(bold=False):
    """Locate a CJK font from the host OS or an explicit user override."""
    override = os.environ.get('PQC_REPORT_FONT_BOLD' if bold else 'PQC_REPORT_FONT')
    candidates = [Path(override)] if override else []
    windows = os.environ.get('WINDIR')
    if windows:
        fonts = Path(windows) / 'Fonts'
        candidates.extend(fonts / name for name in
                          (('msyhbd.ttc', 'simhei.ttf', 'msyh.ttc') if bold else
                           ('msyh.ttc', 'simsun.ttc', 'simhei.ttf')))
    candidates.extend(Path(name) for name in (
        '/System/Library/Fonts/PingFang.ttc',
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc' if bold else
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc' if bold else
        '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'))
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    raise RuntimeError('No CJK font found. Set PQC_REPORT_FONT to a regular CJK font file; '
                       'PQC_REPORT_FONT_BOLD may select a separate bold font.')


class Figure:
    def __init__(self, height):
        self.height = height
        self.im = Image.new('RGB', (2400, round(height * SCALE)), 'white')
        self.d = ImageDraw.Draw(self.im)

    def font(self, size, bold=False):
        return ImageFont.truetype(report_font_path(bold), round(size * SCALE))

    def text(self, xy, text, size=23, fill=INK, bold=False, anchor='mm'):
        x, y = xy
        self.d.multiline_text((x*SCALE, y*SCALE), text, font=self.font(size, bold),
                              fill=fill, anchor=anchor, align='center', spacing=9*SCALE)

    def box(self, rect, text='', fill=PALE, outline=BLUE, size=23, bold=False):
        self.d.rounded_rectangle(tuple(round(v*SCALE) for v in rect), radius=12*SCALE,
                                 fill=fill, outline=outline, width=2*round(SCALE))
        if text:
            self.text(((rect[0]+rect[2])/2, (rect[1]+rect[3])/2), text, size, bold=bold)

    def line(self, points, fill=BLUE, width=2):
        self.d.line([(round(x*SCALE), round(y*SCALE)) for x,y in points], fill=fill,
                    width=max(1, round(width*SCALE)), joint='curve')

    def arrow(self, points, fill=BLUE, width=2.5):
        self.line(points, fill, width)
        (x1,y1),(x2,y2)=points[-2:]
        a=math.atan2(y2-y1,x2-x1)
        head=[(x2,y2),(x2-12*math.cos(a)+6*math.sin(a),y2-12*math.sin(a)-6*math.cos(a)),
              (x2-12*math.cos(a)-6*math.sin(a),y2-12*math.sin(a)+6*math.cos(a))]
        self.d.polygon([(round(x*SCALE),round(y*SCALE)) for x,y in head], fill=fill)

    def save(self, name):
        p=ASSETS/(name+'.png')
        self.im.save(p, dpi=(400,400))
        return p


def make_figures():
    f=Figure(380)
    f.text((500,35),'ML-KEM 在通信系统中的位置',28,bold=True)
    for x,title,sub in [(20,'KeyGen','接收方生成公私钥'),(365,'Encaps','发送方封装'),(710,'Decaps','接收方解封装')]:
        f.box((x,115,x+270,215), title+'\n'+sub,size=23)
    f.arrow([(290,165),(365,165)]);f.text((327,94),'公钥 ek',20)
    f.arrow([(635,165),(710,165)]);f.text((672,94),'密文 c',20)
    f.line([(155,215),(155,265),(845,265)],TEAL)
    f.arrow([(845,265),(845,215)],TEAL)
    f.text((500,241),'私钥 dk 由接收方保存',20,TEAL)
    f.text((500,331),'有效密文与匹配密钥得到相同的 32 B 共享秘密\n后续对称加密与身份认证由应用协议完成',22)
    f.save('01_kem_context')

    f=Figure(380)
    labels=['固定软件基线\n和官方输入','识别高调用量\nKeccak 服务','HLS 建模与\n局部差分验证','接入 RISC-V\n及 RTL 验证']
    for i,label in enumerate(labels):
        x=20+i*250;f.box((x,35,x+210,125),label,size=22)
        if i<3:f.arrow([(x+210,80),(x+250,80)])
    f.arrow([(875,125),(875,205)])
    labels=['归档原始记录\n与独立复核','同板双固件\n正确性与性能','PS–PL 接入\n与固件装载','Vivado 实现\n资源时序签核']
    for i,label in enumerate(labels):
        x=20+i*250;f.box((x,205,x+210,295),label,fill=GREEN,outline=TEAL,size=22)
        if i<3:f.arrow([(x+250,250),(x+210,250)],TEAL)
    f.text((500,348),'日常使用代表性子集   稳定版本执行完整 145 条验收',22,GRAY)
    f.save('02_design_flow')

    f=Figure(750)
    f.box((20,20,980,160),'',fill='#F5F8FB')
    f.text((120,48),'Zynq PS',23,bold=True)
    f.box((55,76,400,138),'Jupyter 与 Python 驱动',fill='white',size=24)
    f.box((590,76,945,138),'装载固件  调度  官方输出核对',fill='white',size=22)
    f.arrow([(400,107),(590,107)])
    f.text((500,191),'M_AXI_GP0 → AXI Interconnect → AXI4-Lite',23,BLUE)
    f.arrow([(500,215),(500,245)])
    f.box((20,245,980,720),'',fill='#F7FBFA',outline=TEAL)
    f.text((152,275),'Zynq PL  100 MHz',23,bold=True)
    f.box((55,310,945,375),'自写 AXI4-Lite 包装层   复位  运行控制  状态与计数',fill='white',size=24)
    f.box((55,440,420,550),'128 KiB 双口 RAM\n固件  输入窗口  输出窗口  栈',size=23)
    f.box((565,440,945,550),'PicoRV32 RV32IM-fast\n执行完整 ML-KEM 软件流程',size=23)
    f.arrow([(250,375),(250,440)]);f.arrow([(755,375),(755,440)])
    f.arrow([(420,480),(565,480)]);f.arrow([(565,515),(420,515)])
    f.box((55,610,420,685),'HLS Keccak 核\nSHA3-256/512  SHAKE128/256',fill=GREEN,outline=TEAL,size=23)
    f.box((565,610,945,685),'原生总线 MMIO adapter\n控制桥  数据缓冲  上下文',fill=GREEN,outline=TEAL,size=23)
    f.arrow([(755,550),(755,610)],TEAL)
    f.arrow([(565,638),(420,638)],TEAL)
    f.arrow([(420,666),(565,666)],TEAL)
    f.text((492,594),'AXI-Lite 控制 / BRAM 数据',18,GRAY)
    f.save('03_architecture')

    f=Figure(550)
    f.text((500,30),'逐位精确的 Keccak 服务',28,bold=True)
    f.box((20,95,255,210),'输入按 rate 分块\n32-bit little-endian\n吸收与标准填充',size=23)
    f.box((365,75,715,240),'Keccak-f[1600]\n25 × 64-bit lane 完整分区\n单轮结构复用 24 轮',fill=GREEN,outline=TEAL,size=23)
    f.box((805,95,980,210),'输出打包\n最多 4 B/步',size=23)
    f.arrow([(255,150),(365,150)]);f.arrow([(715,150),(805,150)])
    f.text((540,287),'每轮依次执行  θ → ρ → π → χ → ι',24)
    f.box((365,340,715,440),'上下文 26 × 64-bit\n25 lane + 状态元数据',size=23)
    f.arrow([(700,240),(770,240),(770,390),(715,390)],TEAL)
    f.line([(365,390),(300,390),(300,185)],TEAL)
    f.arrow([(300,185),(365,185)],TEAL)
    f.text((186,343),'SHAKE 续取时\n恢复原上下文',21,TEAL)
    f.text((500,491),'HASH 初始化并输出   SQUEEZE 继续输出   CLEAR 清除上下文',22,GRAY)
    f.text((500,531),'跨 rate 块时再次置换   当前块未耗尽时直接继续输出',20,GRAY)
    f.save('04_keccak')

    f=Figure(430)
    f.text((500,28),'两级地址空间与缓冲所有权',28,bold=True)
    f.box((20,85,480,330),'',fill=PALE)
    f.text((250,120),'PS 物理基址 0x40000000',24,bold=True)
    f.text((250,217),'0x00000–0x1FFFF  CPU RAM 128 KiB\n0x10000  输入窗口 8 KiB\n0x12000  输出窗口 8 KiB\n0x20000  控制与状态寄存器',23)
    f.box((520,85,980,330),'',fill=GREEN,outline=TEAL)
    f.text((750,120),'CPU MMIO 基址 0x50010000',24,bold=True)
    f.text((750,217),'0x0000  HLS 控制与状态\n0x1000  输入缓冲 2048 B\n0x2000  输出缓冲 4096 B\n0x3000  上下文 208 B',23)
    f.text((500,383),'PS 仅在 CPU 复位时可写 RAM   复位或完成及失败态可读',23,GRAY)
    f.save('05_memory')

    f=Figure(650)
    cols=[90,355,635,910]
    for x,l in zip(cols,['Jupyter / PS','AXI / RAM','PL RISC-V','Keccak 服务']):
        f.text((x,35),l,23,bold=True);f.line([(x,67),(x,600)],'#C7D5DF',1.5)
    def msg(a,b,y,label,color=BLUE):
        f.arrow([(cols[a],y),(cols[b],y)],color)
        f.text(((cols[a]+cols[b])/2,y-26),label,20,color)
    msg(0,1,115,'复位下写固件并读回 SHA-256')
    msg(0,1,205,'写输入  清输出  释放 CPU')
    msg(1,2,280,'邮箱解析与 ML-KEM API')
    msg(2,3,365,'FIPS 202 provider 请求',TEAL)
    msg(3,2,435,'摘要或续取结果',TEAL)
    msg(2,1,510,'写返回值  输出  周期与状态')
    msg(1,0,590,'读回并逐字节核对官方结果')
    f.save('06_transaction')

    f=Figure(450)
    top=[('冻结输入','ACVP 修订与哈希'),('本地 FIPS 202 差分','hashlib / OpenSSL'),('组件与系统 RTL','真实启动与官方输出')]
    bottom=[('离线复核','原始 ZIP / JSON 保留'),('实板配对测试','同板同 bit 双固件'),('实现签核','路由 时序 DRC')]
    for i,(a,b) in enumerate(top):
        x=20+i*335;f.box((x,45,x+290,170),a+'\n'+b,size=23)
        if i<2:f.arrow([(x+290,108),(x+335,108)])
    f.arrow([(835,170),(835,245)])
    for i,(a,b) in enumerate(bottom):
        x=20+i*335;f.box((x,245,x+290,370),a+'\n'+b,fill=GREEN,outline=TEAL,size=23)
        if i<2:f.arrow([(x+335,308),(x+290,308)],TEAL)
    f.text((500,418),'每层验证回答不同问题   历史结果与新测量独立归档',23,GRAY)
    f.save('07_verification')

    f=Figure(650)
    f.text((500,35),'同一 145 条官方记录的板上时间',28,bold=True)
    f.text((500,78),'越短越好   两种时间口径分别比较',22,GRAY)
    # Common axes avoid exaggerating the reduction.
    x0,x1=270,900
    for tick in [0,5,10,15,20,25]:
        x=x0+(x1-x0)*tick/25
        f.line([(x,130),(x,555)],'#DDE5EB',1)
        f.text((x,594),str(tick),21,GRAY)
    for y,label,value,color in [(160,'CPU-only API',16.9929329,BLUE),(245,'CPU+Keccak API',4.46852439,TEAL),
                                (405,'CPU-only 整批',22.135448812,BLUE),(490,'CPU+Keccak 整批',9.634239928,TEAL)]:
        f.text((250,y+19),label,23,anchor='rm')
        end=x0+(x1-x0)*value/25
        f.d.rectangle((x0*SCALE,y*SCALE,end*SCALE,(y+38)*SCALE),fill=color)
        f.text((end+12,y+19),f'{value:.3f} s',22,color,anchor='lm')
    f.text((650,322),'完整 API  3.8028×',27,TEAL,True)
    f.text((650,561),'PS 整批  2.2976×',27,TEAL,True)
    f.text((500,633),'时间 / s   PL 100 MHz   同板同 bitstream',21,GRAY)
    f.save('08_performance')

    f=Figure(470)
    f.box((20,20,980,390),'',fill='#F4F7F9',outline=GRAY)
    f.text((500,60),'PS 事务墙钟',28,bold=True)
    f.text((500,105),'复位  打包  上传  启动  轮询  读回  官方结果核对',23,GRAY)
    f.box((65,155,935,350),'',fill=PALE)
    f.text((500,192),'RISC-V 完整 API 原始周期',25,bold=True)
    f.text((500,235),'软件运算  驱动  上下文  搬运  等待  清零',23)
    f.box((340,275,660,330),'Keccak 硬件 busy',fill=GREEN,outline=TEAL,size=23)
    f.text((500,437),'示意嵌套边界   面积不表示占比   内层计时不能与外层重复相加',22,GRAY)
    f.save('09_timing_scope')

    f=Figure(560)
    f.text((500,30),'PS–PL 完整实现的器件资源占用',28,bold=True)
    vals=[('LUT',19759,53200,37.14),('FF',17585,106400,16.53),('BRAM36 等效',34.5,140,24.64),('DSP',4,220,1.82)]
    for i,(label,n,total,pct) in enumerate(vals):
        y=120+i*100
        f.text((230,y),label,24,anchor='rm')
        f.d.rectangle((260*SCALE,(y-20)*SCALE,870*SCALE,(y+20)*SCALE),fill='#E9EFF3')
        f.d.rectangle((260*SCALE,(y-20)*SCALE,(260+610*pct/100)*SCALE,(y+20)*SCALE),fill=TEAL if i%2 else BLUE)
        f.text((895,y),f'{pct:.2f}%',23,anchor='lm')
        f.text((565,y+39),f'{n:,} / {total:,}',22,GRAY)
    f.text((500,535),'xc7z020clg400-1   包含 PS–PL 接口  CPU  RAM 与加速器',22,GRAY)
    f.save('10_resources')

    for name,label,details in [
        ('11_photo_panorama','系统全景照片预留','请替换为实际连接场景  包含 PYNQ-Z2 电源 网线与主机'),
        ('12_photo_board','开发板近景照片预留','请替换为板卡正面或斜 45° 照片  保留清晰接口与板卡标识'),
        ('13_screen_jupyter','Jupyter 实际运行界面预留','请替换为双模式通过状态与性能对照截图  不含个人或学校信息')]:
        f=Figure(562.5)
        f.d.rectangle((3*SCALE,3*SCALE,997*SCALE,559*SCALE),fill='#F5F7F9',outline='#B8C6D0',width=2*round(SCALE))
        # Small corner marks keep replacement crop unambiguous.
        for x,y,sx,sy in [(28,28,1,1),(972,28,-1,1),(28,534,1,-1),(972,534,-1,-1)]:
            f.line([(x+40*sx,y),(x,y),(x,y+40*sy)],'#A7B8C5',2)
        f.text((500,240),label,32,GRAY,True)
        f.text((500,300),'16:9  同尺寸替换  '+('非真实界面截图' if name=='13_screen_jupyter' else '非实物照片'),24,GRAY)
        f.text((500,360),details,21,GRAY)
        f.save(name)


def font(run,size=None,bold=None,color='000000',east='宋体'):
    run.font.name='Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
    run.font.color.rgb=RGBColor.from_string(color)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold


class Report:
    def __init__(self,template):
        self.table_count=0
        self.d=Document(template)
        body=self.d._element.body
        for child in list(body):
            if child.tag!=qn('w:sectPr'):body.remove(child)
        styles=self.d.styles
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
            if name not in styles:
                styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st=styles[name]
            st.font.name='Times New Roman'
            st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
            st.font.color.rgb=RGBColor(0,0,0)
            pf=st.paragraph_format
            pf.space_before=Pt(0);pf.space_after=Pt(5)
            pf.line_spacing=1.2
            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st._element.get_or_add_pPr().append(snap)
        styles['Normal'].font.size=Pt(11)
        for name,size in [('Title',20),('Subtitle',11),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
            styles[name].font.size=Pt(size)
            styles[name].font.bold=name!='Subtitle'
            styles[name].paragraph_format.keep_with_next=True
        styles['Heading 1'].paragraph_format.space_after=Pt(10)
        styles['Heading 2'].paragraph_format.space_before=Pt(7)
        styles['Heading 2'].paragraph_format.space_after=Pt(5)
        styles['Caption'].font.size=Pt(9)
        styles['Caption'].font.italic=False
        # Do not change the template section, logo, header relationships or footer.
        self.d.core_properties.title='面向 PYNQ Z2 的 ML KEM 1024 软硬件协同加速系统'
        self.d.core_properties.subject='竞赛技术报告'
        self.d.core_properties.author=''
        self.d.core_properties.last_modified_by=''
        self.d.core_properties.keywords='ML-KEM-1024; Keccak; PYNQ-Z2; FPGA; 实板验证'
        s=self.d.settings._element
        uf=s.find(qn('w:updateFields'))
        if uf is None:uf=OxmlElement('w:updateFields');s.append(uf)
        uf.set(qn('w:val'),'true')

    def page(self,title):
        p=self.d.add_paragraph(title,'Heading 1')
        p.paragraph_format.page_break_before=True
        return p

    def heading(self,text):return self.d.add_paragraph(text,'Heading 2')

    def p(self,text,size=11,indent=True,after=5):
        p=self.d.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(22) if indent else Pt(0)
        p.paragraph_format.space_after=Pt(after)
        p.paragraph_format.widow_control=True
        r=p.add_run(text);font(r,size)
        return p

    def note(self,text):return self.p(text,9.5,False,5)

    def figure(self,name,caption,width=146.5):
        p=self.d.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(3)
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.keep_with_next=True
        pic=p.add_run().add_picture(str(ASSETS/(name+'.png')),width=Mm(width))
        pic._inline.docPr.set('descr',caption)
        p=self.d.add_paragraph(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(7)

    def table(self,headers,rows,widths):
        titles=['作品主要性能指标','系统模块职责','HLS 结构与设计取舍','CPU 与加速器接口约定',
                '分层验证结果','官方数据集修订覆盖','官方操作覆盖','重复 恢复与装载检查',
                '同板性能总量','分操作完整 API 周期','固定代表例的十次重复统计',
                'PS 事务阶段总量','PS–PL 实现时序与路由','工程材料入口']
        self.table_count+=1
        caption=self.d.add_paragraph(f'表 {self.table_count}  '+titles[self.table_count-1],'Caption')
        caption.paragraph_format.keep_with_next=True
        caption.paragraph_format.space_after=Pt(4)
        t=self.d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Mm(w)
        for rowdata in [headers]+rows:
            row=t.rows[0] if rowdata is headers else t.add_row()
            for i,(c,txt,w) in enumerate(zip(row.cells,rowdata,widths)):
                c.width=Mm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=c._tc.get_or_add_tcPr()
                mar=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    item=OxmlElement('w:'+side);item.set(qn('w:w'),'70' if side in ['left','right'] else '65');item.set(qn('w:type'),'dxa');mar.append(item)
                tcPr.append(mar)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
                tcPr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'245477' if rowdata is headers else ('F1F6F9' if len(t.rows)%2==0 else 'FFFFFF'));tcPr.append(shade)
                p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=False
                font(p.add_run(str(txt)),9,bold=rowdata is headers,color='FFFFFF' if rowdata is headers else '000000')
            trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
            if rowdata is headers:trPr.append(OxmlElement('w:tblHeader'))
        spacer=self.d.add_paragraph()
        spacer.paragraph_format.space_after=Pt(3)
        spacer.paragraph_format.line_spacing=Pt(1)
        font(spacer.add_run(''),1)
        return t


def build(template):
    r=Report(template)
    p=r.d.add_paragraph('面向 PYNQ Z2 的 ML KEM 1024\n软硬件协同加速系统','Title')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(14)
    p=r.d.add_paragraph('竞赛技术报告','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r.note('版本日期  2026 年 10 月 24 日    测量日期  2026 年 10 月 23 日')
    r.heading('摘要')
    r.p('面向资源受限 FPGA 上的后量子密码实现，本项目在 PYNQ-Z2 上构建了 ML-KEM-1024（K=4）软硬件协同系统。完整算法运行于 PL 内的 PicoRV32 RV32IM-fast 软核，Keccak/SHAKE 运算通过独立 FIPS 202 provider 卸载至 HLS 加速器。PS 端利用 AXI4-Lite、共享 RAM 和 Jupyter 完成固件装载、测试调度与结果核对。设计保留标准算法流程，将优化重点放在高调用量服务、32-bit 数据打包、上下文续取以及接口与时序实现。')
    r.p('在同一开发板、同一 100 MHz bitstream、128 KiB RAM 和 32 KiB 栈条件下，分别切换 CPU-only 与 CPU+Keccak 固件，两组各通过 145 条公开官方测试记录。完整 API 周期由 1,699,293,290 降至 446,852,439，实测加速 3.8028 倍；计入 PS 上传、轮询、读回和校验后，整批墙钟由 22.135 s 降至 9.634 s，仍取得 2.2976 倍收益。代表操作的 60 次重复和六项限定恢复检查全部通过。完整 PS–PL 实现使用 19,759 LUT、17,585 FF、34.5 BRAM36 等效和 4 DSP，在 100 MHz 约束下完成时序签核。')
    r.p('本项目的贡献在于把局部密码运算优化落实为可测量、可复现的完整 ML-KEM 系统，并区分 API、硬件活动及主机事务的计时边界。现有结果支持资源受限平台上的硬件服务卸载与工程验证价值；比较对象为 PL RISC-V 软件，尚不涵盖 PS ARM 优化软件、板级功耗或产品级安全认证。')
    r.note('关键词：后量子密码；ML-KEM-1024；Keccak；HLS；软硬件协同；实板验证')
    r.heading('报告阅读说明')
    r.p('第一部分说明作品用途与主要结果；第二部分给出模块关系、接口和验证方法；第三部分呈现实板证据、性能与资源；第四部分讨论工程价值、限制和扩展；第五部分列出参考资料。全部测试数据来自已归档记录。照片与界面截图的预留位置均明确标识，不以示意图替代实物成果。',size=10)

    r.page('第一部分 作品概述')
    r.heading('设计背景')
    r.p('FIPS 203 已将 ML-KEM 标准化[1]。面向量子计算对传统公钥密码体系的威胁，边缘平台需要评估后量子算法的计算与存储代价。本项目研究如何在商用 FPGA 上通过可复用硬件服务降低完整 KEM 延迟，同时保留软件可编程性和可审计接口。')
    r.heading('功能与特性')
    r.p('作品实现 ML-KEM-1024 的密钥生成、封装、解封装及密钥检查，支持有效输入、无效密钥和隐式拒绝相关测试。用户在 Jupyter 中选择软件或加速固件，执行代表用例或全量回归；系统返回官方输出比对结果、API 周期、硬件计数和原始 JSON。运行前核对 bitstream、固件与向量身份，固件在 CPU 复位下写入完整 RAM 并读回校验。')
    r.heading('应用领域')
    r.p('ML-KEM 是用于建立共享秘密的密钥封装机制，可作为安全通信握手的密码组件[1]。K=4 表示算法的模块维度，对应 ML-KEM-1024 与 NIST 安全类别 5；公钥、私钥、密文分别为 1568、3168、1568 B，共享秘密为 32 B。该原型面向边缘设备密码协处理研究、可重构安全计算教学和竞赛演示。当前没有实现 TLS、证书认证和会话密钥生命周期，应用边界为密码本原及其调度验证。')
    r.figure('01_kem_context','图 1  ML-KEM 的输入输出关系与应用边界')
    r.heading('主要技术特点')
    r.p('完整 KEM 流程保留在可编程软核中，统一卸载 SHA3-256/512 与 SHAKE128/256；HLS 采用 25-lane 状态分区和单轮复用，保持密码运算逐位精确。BRAM 数据缓冲与 MMIO 控制分离，SHAKE 通过上下文支持续取。PS–PL 层提供运行所有权、复位、状态与性能计数，两个 Tcl 入口分别管理主流程和 Block Design，兼顾可观察性与复现。')

    r.page('作品指标与设计流程')
    r.heading('主要性能指标')
    r.table(['指标','已测结果','测量口径'],[
        ['官方记录验证','两组各 145/145','同板双固件'],
        ['完整 API 加速','3.8028×','原始周期总和之比'],
        ['PS 整批加速','2.2976×','22.135 s → 9.634 s'],
        ['系统时钟','100 MHz','实现签核与板上运行'],
        ['逻辑与存储','19,759 LUT / 34.5 BRAM36','完整 PS–PL 实现'],
        ['重复与恢复','60 次 / 6 项通过','固定输入与限定路径']],[44,47,55.5])
    r.heading('主要创新点')
    r.p('一是面向完整 KEM 工作负载，选择可复用的 Keccak 服务卸载并保留标准软件路径；二是将 32-bit 打包、SHAKE 上下文与软核接口协同设计；三是用同板同 bitstream 双固件与分层计时验证净收益，形成从官方输入到实板原始记录的可追溯链路。上述贡献属于系统集成与工程方法，不主张新的密码算法。')
    r.heading('设计流程')
    r.p('先冻结软件、输入与计时边界，再验证独立 HLS 功能、CPU 接口和官方 RTL 输出；完成资源与时序签核后接入 PS–PL，并在同板条件下切换两种固件。日常迭代使用代表子集，稳定版本执行全量验收，失败记录与历史结果分开保留。')
    r.figure('02_design_flow','图 2  从软件基线到实板配对验收的设计流程')
    r.note('历史路线对照：K=2 的细粒度 BaseMul 方案端到端约 0.9930×。该负结果促使项目转向高调用量服务及接口协同优化；其参数与实现不同，不用作 K=4 的同条件消融。')

    r.page('第二部分 系统组成及功能说明')
    r.heading('整体架构')
    r.p('系统由 PS 主机控制、PL 软件执行与 Keccak 硬件服务三层组成。PS 负责装载、调度和核对；完整 ML-KEM 软件实际运行于 PL 的 PicoRV32 RISC-V 软核[5]。PS 的 AXI 主口先接自写包装层，CPU 再经另一地址空间访问 Keccak adapter，二者是分层控制与计算接口。')
    r.note('术语：PS 是芯片中的 ARM 处理系统；PL 是可配置的逻辑电路区域；HLS 将 C++ 运算描述综合为硬件。加速器在 PL 内配合软核执行，Jupyter 提供操作与展示界面。')
    r.figure('03_architecture','图 3  PS 主机与 PL 软硬件协同系统总体框图')
    r.table(['模块','承担任务','实现形式'],[
        ['PS 主机','固件切换、输入输出核对、展示','Python / Jupyter'],
        ['算法执行','NTT、多项式、采样、压缩及 API','PicoRV32 上的 C 固件'],
        ['密码服务','FIPS 202 哈希与可扩展输出','HLS C++ 生成 RTL'],
        ['接口与存储','AXI、MMIO、RAM、状态和计数','自写 RTL / FPGA BRAM']],[31,66,49.5])

    r.page('Keccak 硬件服务与 HLS 结构')
    r.p('ML-KEM 调用 FIPS 202 中的哈希和可扩展输出函数[2]。本项目将这组函数收敛为共享服务，不复制四个 Keccak 核；软件中的 x4 接口映射为四条独立流的串行调用和分别保存的上下文。每条 64-bit lane 保持完整精度，不进行有损定点化或密码状态截断。')
    r.figure('04_keccak','图 4  HLS 加速器的吸收 置换 输出与上下文关系')
    r.table(['结构选择','具体实现','目的与代价'],[
        ['置换结构','25 lane 分区；单轮复用 24 轮','避免展开全部 24 轮的资源增长'],
        ['输入输出','按 rate 分块；32-bit 打包','减少逐字节访存与变量取模'],
        ['状态续取','HASH / SQUEEZE / CLEAR','支持拒绝采样；需保存恢复上下文'],
        ['容量边界','输入 2048 B；单次输出 4096 B','覆盖当前调用形状；非通用无限长消息']],[27,58,61.5])
    r.p('独立 C 仿真和 Verilog 协同仿真各覆盖 109 笔本地差分事务，包括 74 个 HASH 向量、分段输出、交错上下文、边界、清除和错误调用。优化配置的协同仿真总周期由 1,314,553 降至 75,280，局部比值 17.462×。该计数含生成驱动的控制操作，不是纯置换周期，也不能替代完整 KEM 的 3.8028× 实板结果。',size=10)

    r.page('RTL 接口与存储组织')
    r.figure('05_memory','图 5  PS 物理地址与 CPU 内部 MMIO 的分层映射')
    r.heading('外部 AXI 与 RAM 所有权')
    r.p('PS7 的 M_AXI_GP0 经 AXI Interconnect 接自写 32-bit AXI4-Lite slave[7]。映射窗口为 256 KiB，其中低 128 KiB 为 CPU RAM；输入和输出邮箱各占 8 KiB，属于这片 RAM。PS 仅在 CPU 复位时可写 RAM，在复位或完成、失败态可读，运行期间避免主机改写活动固件。数据以单事务 MMIO 交换，未采用 DMA 或零拷贝。')
    r.heading('CPU 到 HLS 的控制与缓冲')
    r.p('PicoRV32 原生总线访问 MMIO adapter，adapter 将控制操作桥接为 HLS AXI-Lite，将数据访问映射到 BRAM。busy 期间缓冲归 HLS 所有；CPU 的冲突访问不写存储并置错。状态包括 busy、done、error；计数分别记录启动、busy 周期和 32-bit 缓冲读写事务。复位不自动擦除 BRAM，软件必须初始化使用的输入和上下文。')
    r.table(['对象','约定','验证或安全边界'],[
        ['返回值','0 成功；参数与上下文错误独立返回','输出及返回码逐条核对'],
        ['上下文','26 × 64-bit；CPU 以 52 个 word 访问','调用方隔离与清除'],
        ['命令互斥','busy 时禁止重复启动与改缓冲','不代表任意故障均可恢复'],
        ['数据端序','little-endian 32-bit 打包','尾 word 高位按接口规则清零']],[27,65,54.5])
    r.note('标准库状态与 HLS 上下文的布局不同，不能直接强制类型转换。CLEAR 清上下文，不应据此声称全部数据 BRAM 已完成密码级擦除。')

    r.page('软件组织与单次执行流程')
    r.p('算法采用固定提交的 mlkem-native portable C[4]，不改写其密码算法。CPU-only 固件使用软件 FIPS 202；加速固件通过 provider 替换该服务且不链接软件 Keccak。NTT、多项式、采样、压缩和封装解封装控制仍由软核执行。测试使用确定性官方输入，未测量实时熵源。')
    r.figure('06_transaction','图 6  从 Jupyter 调度到官方结果核对的事务顺序')
    r.p('驱动首先检查 bit/hwh、固件与向量哈希，在 CPU 复位时写入完整 128 KiB RAM，再读回核对。每个用例清理邮箱、写入参数和数据后释放 CPU；固件执行 API 并提交返回值、输出和周期。主机轮询结束状态、读回结果并逐字节核对官方期望。SHAKE 续取由 provider 恢复对应流的上下文，再向硬件请求后续字节。')
    r.note('驱动兼容本次板卡的 Python 3.6.5。Notebook 可正常保存，不把用户可编辑的输出单元当成不可变证据；验收以原始 JSON、绑定的驱动版本与离线复核为准。')

    r.page('分层验证与可复现构建')
    r.figure('07_verification','图 7  独立参考 实现签核与实板证据之间的关系')
    r.table(['层级','验证对象','已有证据'],[
        ['标准软件','三个参数集的官方输出','主机共 435 条记录通过'],
        ['独立 HLS','FIPS 202 函数与上下文','109 笔 C / RTL 差分事务'],
        ['CPU 与 RTL','真实服务调用及完整 K4 输出','冻结核心 145 条通过'],
        ['系统实现','资源 路由 时序与约束','PS–PL 实现签核'],
        ['实板','两模式官方输出与性能','各 145 条 60 次重复 6 项恢复'],
        ['离线复核','输入身份 哈希 输出 计数与汇总','INDEPENDENT_PAIRED_CHECK_PASS']],[26,58,62.5])
    r.p('工具基线为 Vivado / Vitis HLS 2024.2[6–8]、RISC-V 工具链和 Python 3。主线入口 flow.tcl 以顶部开关管理 HLS、CPU 测试和独立 PL 阶段；bd_axi.tcl 管理 Block Design、AXI 验证、综合实现与导出。源码路径相对仓库定位，生成目录独立；遇到 AMD 长路径限制时用可配置短构建区，不改写历史证据。')
    r.p('每次验收固定源码、TB、固件、向量、bit/hwh 与配置身份。历史报告保留原始字节，目录整理或新实板测量不会反写旧哈希；发布 manifest 的历史状态与后续实板记录分别保存。这样可以追问每个结论由哪一组输入、哪一个实现和哪一次运行支持。',size=10)

    r.page('第三部分 完成情况及性能参数')
    r.heading('实物与部署成果')
    r.p('已完成商用 PYNQ-Z2 上的系统部署与实板验收。作品不包含自制机械结构或 PCB；硬件成果为 PL 逻辑、PS–PL 接口、实现工程与配套 bit/hwh。下列区域待补实际照片，保持相同宽高比即可替换。',size=10)
    r.figure('11_photo_panorama','图 8  系统全景照片预留  建议包含供电 网络与主机连接')
    r.figure('12_photo_board','图 9  PYNQ-Z2 实物近景照片预留  与图 8 同尺寸')

    r.page('官方记录覆盖与恢复检查')
    r.p('2026 年 10 月 23 日，CPU-only 与 CPU+Keccak 两组各完成 145/145 条不同索引，290 条输出及返回值全部核对一致[10]。145 条是指定公开数据集的记录数，并非 145 次完整 KeyGen→Encaps→Decaps 链路。输入来自固定 ACVP-Server 提交[3]，保留 dataset、revision、vsId、tgId 与 tcId。')
    r.table(['数据来源','记录数','处理原则'],[
        ['FIPS203 keyGen','25','正式修订 单独追踪'],
        ['FIPS203 encapDecap','55','正式修订 单独追踪'],
        ['FIPS203-tr1 encapDecap','65','兼容修订 不替代正式修订'],
        ['合计','145','两模式使用相同输入与期望']],[68,20,58.5])
    r.table(['操作类别','条数','覆盖内容'],[
        ['KeyGen / Encaps','25 / 50','密钥与封装输出'],
        ['Decaps / seed Decaps','20 / 10','含拒绝；后者先重建私钥'],
        ['公钥 / 私钥检查','20 / 20','有效性与返回码']],[56,24,66.5])
    r.heading('重复与恢复的已测范围')
    r.table(['检查项目','执行数量','结果与边界'],[
        ['代表性重复','2 模式 × 3 例 × 10 次','60/60；固定 API 周期一致'],
        ['复位后有效用例','两模式各一次','通过限定复位路径'],
        ['错误邮箱 magic','两模式各一次','错误码 2；未启动 Keccak'],
        ['主机极短超时及恢复','两模式各一次','预期超时；复位后有效例通过'],
        ['完整固件装载及读回','共 8 次 每次 128 KiB','预期与读回 SHA-256 一致']],[48,42,56.5])
    r.note('恢复检查不证明任意非法地址、硬件忙时中断或物理故障可恢复。短时重复也不能替代小时级压力、温度和负载变化测试。公开官方向量通过不等于取得 CAVP 或 FIPS 认证。')

    r.page('同板配对性能与主要结果')
    r.p('板上环境为 PYNQ 2.5[9]、Python 3.6.5 与 Linux 4.19。两组采用同一开发板、bitstream、100 MHz、RV32IM-fast、128 KiB RAM、32 KiB 栈、O3 与相同输入。CPU-only 块先运行，加速块后运行，未随机交错。Keccak 电路始终存在，比较服务启用与否的性能，不提供移除该核后的面积或功耗差值。')
    r.figure('08_performance','图 10  完整 API 时间与包含 PS 操作的整批墙钟')
    r.table(['指标','CPU-only','CPU+Keccak'],[
        ['完整 API 原始周期','1,699,293,290','446,852,439'],
        ['API 时间 100 MHz','16.992933 s','4.468524 s'],
        ['逐例 PS 事务之和','22.128231 s','9.627048 s'],
        ['PS 整批墙钟','22.135449 s','9.634240 s'],
        ['固件装载及读回 单列','2.197310 s','2.170197 s']],[55,45.75,45.75])
    r.note('加速比用两组总量相除：API 为 3.802806×，PS 整批为 2.297581×，逐例 PS 事务之和为 2.298548×。不取逐例加速比的算术平均。固件装载、overlay 装载、JSON 保存与浏览器网络不计入性能比值。')

    r.page('操作差异与短时重复性')
    r.p('不同操作对 FIPS 202 的依赖不同，收益不宜用单一局部核指标概括。下表按完整官方集合求和，包含未加速或主机开销主导的类别，未只选择有利用例。')
    r.table(['操作','条数','软件周期','加速周期','API 倍数'],[
        ['KeyGen','25','329,299,932','78,355,067','4.203'],
        ['Encaps','50','700,468,526','192,648,502','3.636'],
        ['Decaps','20','331,730,752','92,092,680','3.602'],
        ['seed Decaps','10','295,059,840','77,103,770','3.827'],
        ['公钥检查','20','5,148,150','5,148,150','1.000'],
        ['私钥检查','20','37,586,090','1,504,270','24.986']],[29,12,37,37,31.5])
    r.p('公钥检查未调用硬件，API 周期完全相同，其 PS 事务比值为 0.9580×，体现主机固定开销与短时波动；不能宣称所有操作均有提升。seed Decaps 包含根据种子重建私钥，不与普通 Decaps 混为同一任务。')
    r.heading('代表例各重复十次')
    r.table(['模式与索引','API 周期','PS 中位数 ms','PS P95 ms'],[
        ['软件 0 KeyGen','13,046,690','193.987','194.425'],
        ['软件 1 Encaps','13,908,285','175.892','176.121'],
        ['软件 115 Decaps','17,088,701','198.003','198.236'],
        ['加速 0 KeyGen','3,120,941','94.677','95.225'],
        ['加速 1 Encaps','3,841,521','75.316','75.517'],
        ['加速 115 Decaps','4,658,615','73.611','73.783']],[48,36,31.25,31.25])
    r.note('各行十次固定输入的 API 周期均一致。P95 使用最近秩 ceil(0.95n)，n=10 时等于本组最大值；该小样本不支持高置信度尾延迟结论。PS 时间来自板上 Python 计时，非浏览器计时。')

    r.page('计时边界与接口成本解释')
    r.figure('09_timing_scope','图 11  事务 API 与硬件活动的包含关系  图形不表示时间占比')
    r.p('API 由固件 rdcycle 原始差值计量，包含 API 内软件计算、provider、打包搬运、上下文、启动等待与清零；不含 PS 邮箱传输、固件启动、官方结果核对和随机源采集。空计时区间单独保留，不在本报告中混用扣除后的值。')
    r.table(['PS 逐例阶段总量','CPU-only s','CPU+Keccak s'],[
        ['上传 含主机打包','2.641','2.641'],
        ['运行等待 含 API 与轮询','17.470','4.981'],
        ['下载','1.817','1.813'],
        ['完整逐例事务和 含核对','22.128','9.627']],[67,39.75,39.75])
    r.p('上传和下载的绝对成本接近，因此 API 的 3.8028× 收益被压缩为约 2.30× 的 PS 收益。运行等待包含 API，阶段合计与 API 不可重复相加；完整事务还含复位、状态和核对等步骤。批次墙钟比逐例之和略长，差额是批次循环与记录开销。')
    r.p('全量加速组记录到 3,188 次 HLS 启动、1,980,263 个 busy 周期、257,851 次缓冲写与 435,632 次缓冲读；软件组对应计数均为零。读写计数的单位是 32-bit word 事务而非周期。API 减 busy 同时含软件和接口成本，不能把差值命名为搬运开销。计数由板上驱动观测，不是外部仪器或密码学证明。',size=10)

    r.page('实现资源与时序签核')
    r.figure('10_resources','图 12  完整 PS–PL overlay 的资源占用  非独立 HLS 核占用')
    r.table(['实现条件或结果','记录值','解释'],[
        ['器件与时钟','xc7z020clg400-1 / 100 MHz','PYNQ-Z2 目标'],
        ['Setup slack','+0.302 ns','100 MHz 约束下通过'],
        ['Hold slack','+0.026 ns','本候选通过'],
        ['Pulse width slack','+3.750 ns','本候选通过'],
        ['路由完成','29,758 / 29,758 网络','实现报告记录'],
        ['签核阻断 DRC','0','发布候选检查结果']],[51,46,49.5])
    r.p('这些结果表明完整设计能在目标器件实现并以本次约束运行。资源余量为后续模块提供可能性，但正时序裕量不等于增量设计可直接继承签核；加入 DMA、多上下文槽或新算法后仍需重新实现。100 MHz 是已实现并实测的工作点，不能写成通过频率扫描得到的 Fmax。')
    r.note('数据来源为 PS–PL 候选 20261005_psaxi06 的实现报告。独立 HLS IP 和独立 PL 自检的资源与边界不同，本页不与其混算。BRAM36 等效 34.5 可包含 BRAM18 的折算；板级功耗和能效尚未测量。')

    r.page('第四部分 总结')
    r.heading('作品价值与贡献定位')
    r.p('本项目完成了从算法软件、HLS 运算核、软核接口、RTL 验证、实现签核到 PS–PL 实板演示的工程闭环。其直接价值是证明在低成本 FPGA 上，选择覆盖真实调用的服务并处理接口成本，可以让完整 ML-KEM 获得净收益，而非只展示一个脱离系统的快核。两组各 145 条官方记录以及同板配对实验支持这一结论。')
    r.p('对工程应用而言，固定接口与标准软件 provider 的边界便于迁移和替换；对研究而言，局部加速与系统收益之间的差距提供了可检验问题：调用粒度、上下文存放和主机传输如何影响资源与延迟。当前尚未做同器件相关工作对照和逐项结构消融，因此不把集成完整性等同于架构最优，也不宣称优于 PS ARM 的优化实现。')
    r.heading('可扩展之处')
    r.p('优先在冻结验收版本上测量 provider 各阶段，再评估批量缓冲、多上下文驻留、减少轮询或 DMA。新增配置逐项比较延迟、资源和时序，避免并行度提升却加重接口成本。面向部署，应补 PS ARM 优化软件对照、仪器功耗、长时压力与温度测试；面向安全产品，还需真实随机源、密钥生命周期、擦除、侧信道与故障注入评估。K=3 及其他参数的硬件支持需独立验证，不能由 K=4 结果直接外推。')
    r.heading('工程总结与设计反思')
    r.p('早期 K=2 BaseMul 的约 0.9930× 结果提醒我们，减少局部算术周期不一定缩短完整 API。后续选择 Keccak 服务，既扩大了可卸载调用范围，也保留了标准算法流程。这次调整的依据是调用工作负载和实测收益，而不是最高参数本身。')
    r.p('HLS 优化同样需要边界意识。对密码状态进行有损定点化会改变标准输出，本设计只优化数据布局、循环结构和存储访问。曾通过局部测试的核仍要在真实 CPU、BRAM 仲裁、时钟复位和主机协议下重新验证；单独 IP 的时序通过也不能代替系统时序通过。')
    r.p('项目最终把输入、固件、硬件与输出身份连成证据链，保留失败与历史记录，使用子集提高迭代效率，再以全量验收冻结结论。这套方法既降低了重复长仿真的成本，也让性能数字能够被解释和复核。当前系统适合作为研究、教学和竞赛原型；进一步产品化仍需上述安全与运行条件评估。')

    r.page('演示与工程复核入口')
    r.p('现场建议依次展示版本及哈希、两种固件的三个代表操作、性能对照与原始记录入口。完整 145 条的板上运行属于秒级到数十秒量级，远短于逐周期 RTL 仿真；是否现场重跑取决于演示时间与网络条件。既有全量证据可直接复核，不要求为了展示重新执行长仿真。')
    r.figure('13_screen_jupyter','图 13  Jupyter 双模式测试与性能对照的实际截图预留')
    r.table(['材料','仓库内相对入口'],[
        ['主线源码与说明','mainline/k4/README.md'],
        ['HLS 与系统流程','mainline/k4/flow.tcl'],
        ['Block Design 与 AXI 流程','mainline/k4/bd_axi.tcl'],
        ['配对测量部署快照','results/k4/paired_firmware/\nboard_20261023/upload_snapshot/'],
        ['本次实板原始证据','results/k4/paired_firmware/board_20261023/']],[48,98.5])
    r.note('已归档证据的只读复核命令  从仓库根目录运行')
    p=r.p('python mainline/k4/scripts/pynq_overlay.py\n  collect-paired results/k4/paired_firmware/board_20261023\n  --check-existing',9,False)
    for run in p.runs:run.font.name='Consolas'
    r.note('以上三行合并为一条命令。只读复核重新检查身份、输出与汇总，无需板卡或 Vivado。重建请按主线 README 配置工具和 Tcl 开关；CHECK_ONLY 与仿真启动不属于最终 PASS。')

    r.page('第五部分 参考文献')
    refs=[
        '[1] NATIONAL INSTITUTE OF STANDARDS AND TECHNOLOGY. Module-Lattice-Based Key-Encapsulation Mechanism Standard: FIPS 203[S]. Gaithersburg: NIST, 2024. DOI: 10.6028/NIST.FIPS.203.',
        '[2] NATIONAL INSTITUTE OF STANDARDS AND TECHNOLOGY. SHA-3 Standard: Permutation-Based Hash and Extendable-Output Functions: FIPS 202[S]. Gaithersburg: NIST, 2015. DOI: 10.6028/NIST.FIPS.202.',
        '[3] NIST. ACVP-Server[CP/OL]. 固定提交 975de31eb83d87039ec88934fdc47d8c312b892d, 采集日期 2026-09-23. https://github.com/usnistgov/ACVP-Server.',
        '[4] PQ CODE PACKAGE. mlkem-native[CP/OL]. 固定提交 b3ba7b32773e657dd37f6f87bce82528459ad8a4. https://github.com/pq-code-package/mlkem-native.',
        '[5] YOSYSHQ. PicoRV32 A Size-Optimized RISC-V CPU[CP/OL]. https://github.com/YosysHQ/picorv32.',
        '[6] AMD. Vitis High-Level Synthesis User Guide: UG1399[M/OL]. 2024.2. https://docs.amd.com/r/2024.2-English/ug1399-vitis-hls.',
        '[7] AMD. Vivado Design Suite User Guide Designing IP Subsystems Using IP Integrator: UG994[M/OL]. 2024.2. https://docs.amd.com/r/2024.2-English/ug994-vivado-ip-subsystems.',
        '[8] AMD. Vivado Design Suite User Guide Using Constraints: UG903[M/OL]. 2024.2. https://docs.amd.com/r/2024.2-English/ug903-vivado-using-constraints.',
        '[9] PYNQ PROJECT. PYNQ Documentation[EB/OL]. v2.5. https://pynq.readthedocs.io/en/v2.5/.',
        '[10] 本项目实板证据归档. 同板双固件配对验证记录及独立核验[R]. 2026-10-23. results/k4/paired_firmware/board_20261023/summary.md 及 verification.json.'
    ]
    for ref in refs:
        p=r.p(ref,10,False,6)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Mm(7)
        p.paragraph_format.first_line_indent=Mm(-7)
    r.heading('数据与引用说明')
    r.p('标准用于定义算法与接口行为，公开软件用于实现与独立核对，AMD 和 PYNQ 文档用于工具与平台。本文的性能、资源和时序数据均引用本项目实际归档，未以厂商标称值或文献中的局部核结果替代实测。公开向量固定在上述提交，FIPS203 与 FIPS203-tr1 修订分别追踪。')
    r.p('本次原始证据 ZIP 的 SHA-256 为：',10,False)
    p=r.p('830e7fbd94c1e8ca621203514bce866ed\n2c8efffb5ffe214ac19f8159a7d37e6',9,False)
    for run in p.runs:run.font.name='Consolas'
    r.note('哈希换行为排版需要，核对时连接为 64 个十六进制字符。原 ZIP、42 个 JSON、上传绑定快照与离线复核脚本均独立保留；后续补照片不改变实验原始数据。')
    out=HERE/'MLKEM1024_竞赛技术报告.docx'
    r.d.save(out)
    # Restore opaque template parts byte-for-byte after python-docx serialization.
    # The body, selected body styles, settings, core metadata and new media are editable.
    with zipfile.ZipFile(template) as source,zipfile.ZipFile(out) as final:
        edited={n:final.read(n) for n in final.namelist()}
        editable={'word/document.xml','word/styles.xml','word/settings.xml',
                  'word/_rels/document.xml.rels','docProps/core.xml','[Content_Types].xml'}
        preserve=[n for n in source.namelist() if n not in editable]
        for n in preserve:edited[n]=source.read(n)
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in edited.items():z.writestr(n,b)
    stats={'docx':str(out),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),
           'reference_sha256':hashlib.sha256(template.read_bytes()).hexdigest(),
           'figures':len(list(ASSETS.glob('*.png'))),'tables':len(r.d.tables),
           'explicit_photo_slots':3,'preserved_template_parts':preserve}
    (ROOT/'build/competition_report/authoring.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(stats,ensure_ascii=False,indent=2))


if __name__=='__main__':
    ap=argparse.ArgumentParser()
    ap.add_argument('--template',type=Path,
                    default=HERE/'reference_template.docx',
                    help='DOCX competition template (default: bundled reference_template.docx)')
    args=ap.parse_args()
    make_figures()
    build(args.template)
